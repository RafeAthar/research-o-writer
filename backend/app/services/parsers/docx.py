from __future__ import annotations

import re
from io import BytesIO

from docx import Document

from app.services.parsers.types import ParsedDocument, ParsedParagraph, StructureNode

_HEADING_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


def parse_docx(data: bytes) -> ParsedDocument:
    doc = Document(BytesIO(data))

    core = doc.core_properties
    title = core.title or None
    authors = [a.strip() for a in (core.author or "").split(";") if a.strip()] or (
        [core.author.strip()] if core.author and core.author.strip() else []
    )
    language = (core.language or "").split("-")[0] or None

    paragraphs: list[ParsedParagraph] = []
    structure_root: list[StructureNode] = []
    stack: list[StructureNode] = []
    chapter_path: list[str] = []
    full_parts: list[str] = []
    char_cursor = 0
    para_idx = 0
    word_count = 0

    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue

        style = (p.style.name or "") if p.style else ""
        m = _HEADING_RE.match(style)
        if m:
            level = int(m.group(1))
            node = StructureNode(title=text, depth=level, chapter_path=[])
            while stack and stack[-1].depth >= level:
                stack.pop()
            if stack:
                parent = stack[-1]
                node.chapter_path = [*parent.chapter_path, text]
                node.order_in_parent = len(parent.children)
                parent.children.append(node)
            else:
                node.chapter_path = [text]
                node.order_in_parent = len(structure_root)
                structure_root.append(node)
            stack.append(node)
            chapter_path = list(node.chapter_path)
            continue

        start = char_cursor
        end = start + len(text)
        paragraphs.append(
            ParsedParagraph(
                text=text,
                chapter_path=list(chapter_path),
                paragraph_index=para_idx,
                page_start=None,
                page_end=None,
                char_start=start,
                char_end=end,
            )
        )
        full_parts.append(text)
        char_cursor = end + 2
        para_idx += 1
        word_count += len(text.split())

    full_text = "\n\n".join(full_parts)
    return ParsedDocument(
        title=title,
        authors=authors,
        language=language,
        page_count=None,
        word_count=word_count,
        paragraphs=paragraphs,
        structure=structure_root,
        full_text=full_text,
    )
